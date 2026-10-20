"""Makes the sample files the reader tests read, one per extension sift reads, all with the same content.

The samples are committed, so the tests need none of this. To make them again (after changing the content):

    pip install python-docx python-pptx openpyxl odfpy fpdf2
    python tests/samples/make_samples.py
"""

import datetime
import json
import zipfile
from pathlib import Path

HERE = Path(__file__).parent
TITLE, AUTHOR, KEYWORDS = "Pigment mixing", "Ada Example", "paint, colour"
CREATED = datetime.datetime(2024, 5, 1, 12, 0, 0)
HEADING = "Kubelka-Munk theory"
PARAGRAPH = "Kubelka-Munk theory predicts the colour of mixed paints from how each pigment absorbs and scatters light."


def sentence(fmt):
    return f"This sample is the {fmt} file."


def pdf(path):
    from fpdf import FPDF

    doc = FPDF()
    doc.set_title(TITLE)
    doc.set_author(AUTHOR)
    doc.set_keywords(KEYWORDS)
    doc.set_creation_date(CREATED.replace(tzinfo=datetime.UTC))
    doc.add_page()
    doc.set_font("helvetica", "B", 20)
    doc.cell(text=HEADING, new_x="LMARGIN", new_y="NEXT")
    doc.set_font("helvetica", size=11)
    for text in (PARAGRAPH, sentence("PDF")):
        doc.multi_cell(0, text=text, new_x="LMARGIN", new_y="NEXT")
    doc.output(str(path))


def docx(path):
    import docx

    doc = docx.Document()
    props = doc.core_properties
    props.title, props.author, props.keywords, props.created = TITLE, AUTHOR, KEYWORDS, CREATED
    doc.add_heading(HEADING, level=1)
    doc.add_paragraph(PARAGRAPH)
    doc.add_paragraph(sentence("DOCX"))
    doc.save(path)


def pptx(path):
    import pptx

    deck = pptx.Presentation()
    props = deck.core_properties
    props.title, props.author, props.keywords, props.created = TITLE, AUTHOR, KEYWORDS, CREATED
    slide = deck.slides.add_slide(deck.slide_layouts[1])  # title and content
    slide.shapes.title.text = HEADING
    slide.placeholders[1].text = PARAGRAPH
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "About this sample"
    slide.placeholders[1].text = sentence("PPTX")
    deck.save(path)


def xlsx(path):
    import openpyxl

    book = openpyxl.Workbook()
    props = book.properties
    props.title, props.creator, props.keywords, props.created = TITLE, AUTHOR, KEYWORDS, CREATED
    sheet = book.active
    sheet.title = HEADING
    for row in (
        ["Pigment", "Colour index", "Note"],
        ["Titanium white", "PW6", PARAGRAPH],
        ["Ultramarine", "PB29", sentence("XLSX")],
    ):
        sheet.append(row)
    book.save(path)


def _odf(doc, path):
    from odf.dc import Creator, Title
    from odf.meta import CreationDate, Keyword

    doc.meta.addElement(Title(text=TITLE))
    doc.meta.addElement(Creator(text=AUTHOR))
    doc.meta.addElement(Keyword(text=KEYWORDS))
    doc.meta.addElement(CreationDate(text=CREATED.isoformat()))
    doc.save(str(path))


def odt(path):
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P

    doc = OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text=HEADING))
    doc.text.addElement(P(text=PARAGRAPH))
    doc.text.addElement(P(text=sentence("ODT")))
    _odf(doc, path)


def odp(path):
    from odf.draw import Frame, Page, TextBox
    from odf.opendocument import OpenDocumentPresentation
    from odf.style import MasterPage, PageLayout
    from odf.text import P

    doc = OpenDocumentPresentation()
    layout = PageLayout(name="layout")
    doc.automaticstyles.addElement(layout)
    master = MasterPage(name="master", pagelayoutname=layout)
    doc.masterstyles.addElement(master)
    for title, text in ((HEADING, PARAGRAPH), ("About this sample", sentence("ODP"))):
        page = Page(masterpagename=master)
        doc.presentation.addElement(page)
        for line in (title, text):
            frame = Frame(width="20cm", height="3cm", x="1cm", y="1cm")
            page.addElement(frame)
            box = TextBox()
            frame.addElement(box)
            box.addElement(P(text=line))
    _odf(doc, path)


def ods(path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    table = Table(name=HEADING)
    for row in (
        ["Pigment", "Colour index", "Note"],
        ["Titanium white", "PW6", PARAGRAPH],
        ["Ultramarine", "PB29", sentence("ODS")],
    ):
        tr = TableRow()
        for value in row:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            tr.addElement(cell)
        table.addElement(tr)
    doc.spreadsheet.addElement(table)
    _odf(doc, path)


def rtf(path):
    path.write_bytes(
        (
            r"{\rtf1\ansi\deff0{\fonttbl{\f0 Helvetica;}}"
            rf"{{\info{{\title {TITLE}}}{{\author {AUTHOR}}}{{\keywords {KEYWORDS}}}"
            rf"{{\creatim\yr{CREATED.year}\mo{CREATED.month}\dy{CREATED.day}\hr{CREATED.hour}\min{CREATED.minute}}}}}"
            rf"{{\pard\b\fs40 {HEADING}\par}}"
            rf"{{\pard\fs22 {PARAGRAPH}\par}}"
            rf"{{\pard\fs22 {sentence('RTF')}\par}}}}"
            "\n"
        ).encode("ascii")
    )


def html(fmt):
    return (
        '<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
        f'<title>{TITLE}</title>\n<meta name="author" content="{AUTHOR}">\n'
        f'<meta name="keywords" content="{KEYWORDS}">\n</head>\n<body>\n'
        f"<h1>{HEADING}</h1>\n<p>{PARAGRAPH}</p>\n<p>{sentence(fmt)}</p>\n</body>\n</html>\n"
    )


def epub(path):
    """An EPUB 3 book: a zip with the mimetype first and uncompressed, a container, a package and one chapter."""
    container = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
        '<rootfiles><rootfile full-path="book.opf" media-type="application/oebps-package+xml"/></rootfiles>\n'
        "</container>\n"
    )
    package = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<package version="3.0" unique-identifier="id" xmlns="http://www.idpf.org/2007/opf">\n'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        f'<dc:identifier id="id">sift-sample</dc:identifier><dc:title>{TITLE}</dc:title>\n'
        f"<dc:creator>{AUTHOR}</dc:creator><dc:language>en</dc:language>\n"
        f"<dc:date>{CREATED.date().isoformat()}</dc:date>\n"
        f'<meta property="dcterms:modified">{CREATED.isoformat()}Z</meta>\n'
        "</metadata>\n"
        '<manifest><item id="ch1" href="chapter.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest>\n'
        '<spine><itemref idref="ch1"/></spine>\n'
        "</package>\n"
    )
    nav = (
        '<?xml version="1.0" encoding="utf-8"?>\n<html xmlns="http://www.w3.org/1999/xhtml" '
        'xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body>'
        f'<nav epub:type="toc"><ol><li><a href="chapter.xhtml">{HEADING}</a></li></ol></nav></body></html>\n'
    )
    chapter = html("EPUB").replace(
        "<!doctype html>\n<html", '<?xml version="1.0" encoding="utf-8"?>\n<html xmlns="http://www.w3.org/1999/xhtml"'
    )
    chapter = chapter.replace('<meta charset="utf-8">', '<meta charset="utf-8"/>').replace('">\n<meta', '"/>\n<meta')
    chapter = chapter.replace(f'content="{KEYWORDS}">', f'content="{KEYWORDS}"/>')
    stamp = CREATED.timetuple()[:6]
    with zipfile.ZipFile(path, "w") as book:
        book.writestr(zipfile.ZipInfo("mimetype", stamp), "application/epub+zip", zipfile.ZIP_STORED)
        for name, text in (
            ("META-INF/container.xml", container),
            ("book.opf", package),
            ("nav.xhtml", nav),
            ("chapter.xhtml", chapter),
        ):
            book.writestr(zipfile.ZipInfo(name, stamp), text, zipfile.ZIP_DEFLATED)


def blank_pdf(path):
    """A page without text, as a scan is to sift (it runs no OCR)."""
    from fpdf import FPDF

    doc = FPDF()
    doc.add_page()
    doc.output(str(path))


def text(fmt):
    files = {
        "md": f'---\ntitle: "{TITLE}"\nauthors: ["{AUTHOR}"]\n---\n\n# {HEADING}\n\n{PARAGRAPH}\n\n{sentence("MD")}\n',
        "txt": f"{HEADING}\n\n{PARAGRAPH}\n\n{sentence('TXT')}\n",
    }
    files["markdown"] = files["md"].replace("MD", "MARKDOWN")
    record = {"id": "pigments", "title": TITLE, "authors": [AUTHOR], "text": f"{PARAGRAPH}\n\n{sentence('JSONL')}"}
    files["jsonl"] = json.dumps(record) + "\n"
    files["ndjson"] = files["jsonl"].replace("JSONL", "NDJSON")
    return files[fmt]


def main():
    made = {"pdf": pdf, "docx": docx, "pptx": pptx, "xlsx": xlsx, "odt": odt, "odp": odp, "ods": ods, "rtf": rtf}
    for ext, make in {**made, "epub": epub}.items():
        make(HERE / f"sample.{ext}")
    for ext in ("html", "htm"):
        (HERE / f"sample.{ext}").write_bytes(html(ext.upper()).encode())
    for ext in ("md", "markdown", "txt", "jsonl", "ndjson"):
        (HERE / f"sample.{ext}").write_bytes(text(ext).encode())
    blank_pdf(HERE / "blank.pdf")
    print("samples written to", HERE)


if __name__ == "__main__":
    main()
